from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
EXTRACT_ROOT = ROOT / 'extract' / 'M268 DATA DASHBOARD'
ONEDRIVE_ROOT = ROOT.parents[2]


def source_file(filename):
    live_file = ONEDRIVE_ROOT / filename
    return live_file if live_file.exists() else EXTRACT_ROOT / filename


WORKBOOK = source_file('R Dashboard M268 EX.xlsx')
BNPL_WORKBOOK = source_file('R DATA BNPL M268.xlsx')
LOB_WORKBOOK = source_file('R LOB & Data Base M268 EX.xlsx')
ROSTER_WORKBOOK = source_file('R Roster M268.xlsx')
OUTPUT = ROOT / 'dashboard_data.js'

INCENTIVE_PER_UNIT = {
    'Mac': 30000,
    'iPad': 10000,
    'iPhone': 15000,
    'Apple Watch': 10000,
    'AirPods': 10000,
}


def to_rp(value):
    if value is None:
        return 'Rp 0'
    if isinstance(value, (int, float)):
        return f'Rp {value:,.0f}'
    return str(value)


def parse_dashboard():
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb['Dashboard']
    rekap = wb['M268 REKAP']
    spw = wb['SPW']
    soh = wb['SOH']
    bnpl_wb = load_workbook(BNPL_WORKBOOK, data_only=True, read_only=True)
    bnpl_ws = bnpl_wb['R DATA BNPL']
    lob_wb = load_workbook(LOB_WORKBOOK, data_only=True, read_only=True)
    lob_master = lob_wb['Master']
    roster_wb = load_workbook(ROSTER_WORKBOOK, data_only=True, read_only=True)
    roster_master = roster_wb['MASTER MPP M268']
    roster_schedule = roster_wb['ROSTER']
    database_wb = load_workbook(source_file('R DATA BASE M268.xlsx'), data_only=True, read_only=True)
    database_ws = database_wb['R DATA BASE M268']
    rekap_rows = list(rekap.iter_rows(min_row=5, max_row=15, values_only=True))
    target_row = rekap_rows[1]
    achievement_row = rekap_rows[2]
    total_device_row = rekap_rows[10]
    stock_columns = {'iPhone': (2, 4, 3, 5), 'iPad': (9, 11, 10, 12), 'Mac': (16, 18, 17, 19), 'Apple Watch': (23, 25, 24, 26)}
    stock_updated_at = datetime.fromtimestamp(WORKBOOK.stat().st_mtime).astimezone().isoformat()
    stock = {}
    for label, (article_col, qty_col, description_col, price_col) in stock_columns.items():
        items = []
        for row in soh.iter_rows(min_row=6, values_only=True):
            article, qty, description, price = row[article_col], row[qty_col], row[description_col], row[price_col]
            if article and description:
                unit_price = float(str(price or 0).replace(',', '').replace(' ', '') or 0)
                items.append({'article': str(article).strip(), 'description': str(description).strip(), 'qty': int(qty or 0), 'price': unit_price})
        stock[label] = {'qty': sum(item['qty'] for item in items), 'amount': sum(item['qty'] * item['price'] for item in items), 'items': items}

    current_month = '2026-09'
    product_lob_by_article = {}
    for row in lob_master.iter_rows(min_row=2, values_only=True):
        article, description, lob = row[0], row[1], row[2]
        if article and lob:
            product_lob_by_article[str(article).strip().upper()] = str(lob).strip().upper()
    sales = []
    for row in rekap.iter_rows(min_row=6, max_row=11, values_only=True):
        name = row[27]
        if not isinstance(name, str) or not name.strip():
            continue
        lob_units = {
            'Mac': int(row[16] or 0),
            'iPad': int(row[17] or 0),
            'iPhone': int(row[18] or 0),
            'Apple Watch': int(row[19] or 0),
            'AirPods': int(row[20] or 0),
        }
        incentive = sum(lob_units[lob] * amount for lob, amount in INCENTIVE_PER_UNIT.items())
        sales.append({
            'name': name.strip(),
            'lob': lob_units,
            'totalDevice': int(row[21] or 0),
            'qoala': int(row[22] or 0),
            'provider': int(row[23] or 0),
            'achievement': float(row[28] or 0),
            'target': float(row[29] or 0),
            'revenue': float(row[30] or 0),
            'variance': float(row[31] or 0),
            'deviceTarget': float(row[32] or 0),
            'deviceAchievement': float(row[33] or 0),
            'accessoriesTarget': float(row[35] or 0),
            'accessoriesAchievement': float(row[36] or 0),
            'vasTarget': float(row[38] or 0),
            'vasAchievement': float(row[39] or 0),
            'estimatedIncentive': incentive,
            'products': [],
        })

    sales_by_name = {sale['name'].casefold(): sale for sale in sales}
    roster = []
    for row in roster_master.iter_rows(min_row=6, values_only=True):
        name, nik, position, join_date = row[2], row[3], row[4], row[5]
        if not name:
            continue
        roster.append({
            'name': str(name).strip(),
            'nik': str(nik).strip() if nik else '',
            'position': str(position).strip() if position else '',
            'joinDate': join_date.strftime('%Y-%m-%d') if hasattr(join_date, 'strftime') else str(join_date or ''),
        })
    shift_names = {'D047': 'Morning', 'D103': 'Afternoon', 'D083': 'Middle', 'OFF': 'OFF', 'X&C': 'Leave'}
    schedule_dates = []
    schedule_header = list(roster_schedule.iter_rows(min_row=6, max_row=6, values_only=True))[0]
    for index, value in enumerate(schedule_header):
        if index >= 6 and hasattr(value, 'strftime'):
            schedule_dates.append((index, value.strftime('%Y-%m-%d')))
    schedule_by_name = {}
    for row in roster_schedule.iter_rows(min_row=8, values_only=True):
        name = row[2] if len(row) > 2 else None
        if not name:
            continue
        entries = []
        for index, date_text in schedule_dates:
            code = str(row[index] or '').strip().upper()
            if code:
                entries.append({'date': date_text, 'code': code, 'label': shift_names.get(code, code)})
        schedule_by_name[str(name).strip().casefold()] = entries
    for member in roster:
        member['schedule'] = schedule_by_name.get(member['name'].casefold(), [])
    database = []
    for row in database_ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        start_time, completion_time, customer, phone, product = row[1:6]
        database.append({
            'id': str(row[0] or ''),
            'startTime': start_time.strftime('%Y-%m-%d %H:%M') if hasattr(start_time, 'strftime') else str(start_time or ''),
            'completionTime': completion_time.strftime('%Y-%m-%d %H:%M') if hasattr(completion_time, 'strftime') else str(completion_time or ''),
            'customer': str(customer or '').strip(),
            'phone': str(phone or '').replace('\u202a', '').replace('\u202c', '').strip(),
            'product': str(product or '').strip(),
        })
    all_products = []
    accessory_qty = 0
    vas_qty = 0
    daily_sales = {}
    daily_qoala = {}
    daily_provider = {}
    service_totals = {}
    for row in spw.iter_rows(min_row=2, values_only=True):
        raw_date, raw_name, article, description, category, qty, amount, brand = row[0], row[2], row[4], row[5], row[6], row[7], row[8], row[9]
        if not raw_date or not raw_name or not description:
            continue
        if hasattr(raw_date, 'strftime'):
            date_text = raw_date.strftime('%Y-%m-%d')
        else:
            raw_date_text = str(raw_date)
            parts = raw_date_text[:10].split('-')
            date_text = f'{parts[2]}-{parts[1]}-{parts[0]}' if len(parts) == 3 and len(parts[0]) == 2 else raw_date_text[:10]
        if not date_text.startswith(current_month):
            continue
        name = str(raw_name).strip()
        transaction_qty = int(qty or 0)
        transaction_amount = float(amount or 0)
        category_name = str(category or '').strip().upper()
        if category_name in ('PHONE', 'TABLETS', 'LAPTOPS', 'WATCH'):
            daily_category = 'device'
        elif category_name == 'SOFTWARE':
            daily_category = 'vas'
        else:
            daily_category = 'accessories'
        if transaction_amount:
            daily_sales.setdefault(date_text, {'qty': 0, 'amount': 0, 'sales': {}})
            daily_sales[date_text]['qty'] += transaction_qty
            daily_sales[date_text]['amount'] += transaction_amount
            daily_sales[date_text]['sales'].setdefault(name, {
                'qty': 0, 'amount': 0, 'achievement': 0,
                'deviceAchievement': 0, 'accessoriesAchievement': 0,
                'vasAchievement': 0
            })
            daily_sale = daily_sales[date_text]['sales'][name]
            daily_sale['qty'] += transaction_qty
            daily_sale['amount'] += transaction_amount
            daily_sale['achievement'] += transaction_amount
            daily_sale[f'{daily_category}Achievement'] += transaction_amount
        brand_name = str(brand or '').strip().upper()
        if brand_name in ('KOALA', 'QOALA'):
            daily_qoala.setdefault(date_text, {'qty': 0, 'amount': 0, 'sales': {}})
            daily_qoala[date_text]['qty'] += transaction_qty
            daily_qoala[date_text]['amount'] += transaction_amount
            daily_qoala[date_text]['sales'].setdefault(name, {'qty': 0, 'amount': 0})
            daily_qoala[date_text]['sales'][name]['qty'] += transaction_qty
            daily_qoala[date_text]['sales'][name]['amount'] += transaction_amount
        if brand_name in ('TELKOMSEL', 'INDOSAT', 'XL', 'XXL'):
            daily_provider.setdefault(date_text, {'qty': 0, 'amount': 0, 'sales': {}})
            daily_provider[date_text]['qty'] += transaction_qty
            daily_provider[date_text]['amount'] += transaction_amount
            daily_provider[date_text]['sales'].setdefault(name, {'qty': 0, 'amount': 0})
            daily_provider[date_text]['sales'][name]['qty'] += transaction_qty
            daily_provider[date_text]['sales'][name]['amount'] += transaction_amount
        mapped_lob = product_lob_by_article.get(str(article).strip().upper())
        if brand_name == 'TELKOMSEL':
            vas_qty += int(qty or 0)
            service_totals.setdefault('Telkomsel', {'qty': 0, 'amount': 0})
            service_totals['Telkomsel']['qty'] += int(qty or 0)
            service_totals['Telkomsel']['amount'] += float(amount or 0)
            continue
        text = f'{description} {category or ""}'.lower()
        if mapped_lob == 'IPHONE' or 'iphone' in text:
            unit = 'iPhone'
        elif mapped_lob == 'IPAD' or 'ipad' in text:
            unit = 'iPad'
        elif mapped_lob == 'MAC' or 'mac' in text or 'macbook' in text or 'mba ' in text:
            unit = 'Mac'
        elif mapped_lob in ('AW', 'APPLE WATCH') or 'watch' in text:
            unit = 'Apple Watch'
        elif mapped_lob == 'AIRPODS' or 'airpods' in text:
            unit = 'AirPods'
        else:
            accessory_qty += int(qty or 0)
            continue
        if brand_name != 'APPLE':
            continue
        sale = sales_by_name.get(name.casefold())
        if not sale:
            continue
        product = {
            'salesName': name,
            'unit': unit,
            'description': str(description).strip(),
            'qty': int(qty or 0),
            'amount': float(amount or 0),
            'date': date_text,
        }
        sale['products'].append(product)
        all_products.append(product)

    daily_bnpl = {}
    provider_totals = {}
    promoter_totals = {}
    current_provider_totals = {}
    period_totals = {}
    for row in bnpl_ws.iter_rows(min_row=2, values_only=True):
        sale_date, provider, promoter, amount, qty = row[4], row[3], row[5], row[6], row[7]
        if not sale_date or not provider:
            continue
        date_key = sale_date.strftime('%Y-%m-%d') if hasattr(sale_date, 'strftime') else str(sale_date)[:10]
        month_key = date_key[:7]
        amount = float(amount or 0)
        qty = float(qty or 1)
        provider = str(provider).strip().upper()
        promoter = str(promoter or 'Tidak diketahui').strip()
        provider_totals.setdefault(provider, {'amount': 0, 'qty': 0})
        provider_totals[provider]['amount'] += amount
        provider_totals[provider]['qty'] += qty
        promoter_totals.setdefault(promoter, {'amount': 0, 'qty': 0})
        promoter_totals[promoter]['amount'] += amount
        promoter_totals[promoter]['qty'] += qty
        if month_key == current_month:
            current_provider_totals.setdefault(provider, {'amount': 0, 'qty': 0})
            current_provider_totals[provider]['amount'] += amount
            current_provider_totals[provider]['qty'] += qty
        period_totals.setdefault(month_key, {'amount': 0, 'qty': 0})
        period_totals[month_key]['amount'] += amount
        period_totals[month_key]['qty'] += qty
        daily_bnpl.setdefault(date_key, {'amount': 0, 'qty': 0, 'providers': {}})
        daily_bnpl[date_key]['amount'] += amount
        daily_bnpl[date_key]['qty'] += qty
        daily_bnpl[date_key]['providers'].setdefault(provider, {'amount': 0, 'qty': 0})
        daily_bnpl[date_key]['providers'][provider]['amount'] += amount
        daily_bnpl[date_key]['providers'][provider]['qty'] += qty

    bnpl = [
        {'label': provider.title(), 'value': round(values['amount']), 'qty': round(values['qty']), 'color': color}
        for (provider, values), color in zip(
            sorted(
                current_provider_totals.items(),
                key=lambda item: item[1]['amount'],
                reverse=True,
            ),
            ['#1d4ed8', '#2563eb', '#22c55e', '#f59e0b', '#ef4444', '#8b5cf6']
        )
    ]

    mading_units = []
    official_unit_qty = {
        'Mac': int(total_device_row[16] or 0),
        'iPad': int(total_device_row[17] or 0),
        'iPhone': int(total_device_row[18] or 0),
        'Apple Watch': int(total_device_row[19] or 0),
        'AirPods': int(total_device_row[20] or 0),
    }
    for unit in ('Mac', 'iPad', 'iPhone', 'Apple Watch', 'AirPods'):
        unit_products = [product for product in all_products if product['unit'] == unit]
        mading_units.append({
            'label': unit,
            'qty': official_unit_qty[unit],
            'amount': sum(product['amount'] for product in unit_products),
        })
    service_totals['Qoala'] = {'qty': int(total_device_row[22] or 0), 'amount': float(achievement_row[7] or 0)}
    for label in ('Indosat', 'XXL'):
        service_totals.setdefault(label, {'qty': 0, 'amount': 0})

    summary = {
        'title': 'M268 Executive Dashboard',
        'subtitle': 'September 2026 • Performance & Growth Tracking',
        'metrics': [
            {
                'label': 'Total Achievement',
                'value': 'Rp 2.719.574.700',
                'detail': 'EST: Rp 3.547.271.348 (109%)',
                'accent': '#2f6fed'
            },
            {
                'label': 'Point Store',
                'value': '77',
                'detail': 'Store performance score',
                'accent': '#20b2a6'
            },
            {
                'label': 'Growth Trend',
                'value': '+134%',
                'detail': 'Month-over-month growth',
                'accent': '#16a34a'
            },
            {
                'label': 'Time Gone / Sisa Hari',
                'value': '73% (8 Hari Sisa)',
                'detail': 'Execution pace',
                'accent': '#f59e0b'
            },
            {
                'label': 'Tim Sales',
                'value': '5 Staff',
                'detail': 'Current sales team',
                'accent': '#ec4899'
            },
            {
                'label': 'Daily Target',
                'value': 'Rp 58.936.144',
                'detail': 'Store daily target',
                'accent': '#8b5cf6'
            },
        ],
        'segments': [
            {'label': 'Device', 'value': int(total_device_row[21] or 0), 'color': '#2f6fed'},
            {'label': 'Accessories', 'value': accessory_qty, 'color': '#60a5fa'},
            {'label': 'VAS', 'value': vas_qty, 'color': '#a78bfa'},
            {'label': 'BNPL', 'value': round(sum(item['qty'] for item in bnpl)), 'color': '#14b8a6'},
        ],
        'lob': [
            {'label': 'Mac', 'value': 20, 'color': '#2563eb'},
            {'label': 'iPad', 'value': 27, 'color': '#3b82f6'},
            {'label': 'iPhone', 'value': 149, 'color': '#f59e0b'},
            {'label': 'Apple Watch', 'value': 23, 'color': '#10b981'},
            {'label': 'AirPods', 'value': 18, 'color': '#8b5cf6'},
        ],
        'bnpl': bnpl,
        'promoters': [
            {'label': promoter, 'value': round(values['amount']), 'qty': round(values['qty'])}
            for promoter, values in sorted(promoter_totals.items(), key=lambda item: item[1]['amount'], reverse=True)
        ],
        'mading': {
            'categories': [
                {'label': 'Total', 'target': float(target_row[3] or 0), 'achievement': float(achievement_row[3] or 0)},
                {'label': 'Device', 'target': float(target_row[4] or 0), 'achievement': float(achievement_row[4] or 0)},
                {'label': 'Accessories', 'target': float(target_row[5] or 0), 'achievement': float(achievement_row[5] or 0)},
                {'label': 'VAS', 'target': float(target_row[6] or 0), 'achievement': float(achievement_row[6] or 0)},
            ],
            'units': mading_units,
            'totalPrice': float(achievement_row[3] or 0),
            'services': [{'label': label, **values} for label, values in service_totals.items()],
        },
        'notes': [
            'Data source: Excel workbook M268 Dashboard and M268 Rekap.',
            'Store target pacing is estimated at 109% with 8 working days left.',
            'Primary growth remains driven by iPhone and accessories contribution.',
        ],
        'incentivePerUnit': INCENTIVE_PER_UNIT,
        'sales': sales,
        'dailyBnpl': daily_bnpl,
        'dailySales': daily_sales,
        'dailyQoala': daily_qoala,
        'dailyProvider': daily_provider,
        'periodTotals': period_totals,
        'currentMonth': current_month,
        'products': all_products,
        'roster': roster,
        'stock': stock,
        'stockUpdatedAt': stock_updated_at,
        'database': database,
    }

    return summary


def main():
    data = parse_dashboard()
    output = f"window.DASHBOARD_DATA = {json.dumps(data, ensure_ascii=False, indent=2)};\n"
    OUTPUT.write_text(output, encoding='utf-8')
    print(f'Created {OUTPUT.name} from {WORKBOOK.name}')


if __name__ == '__main__':
    main()
